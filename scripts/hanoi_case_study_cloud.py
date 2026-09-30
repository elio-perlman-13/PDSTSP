#!/usr/bin/env python3
"""GitHub Actions helpers for the 100-customer Hanoi case study."""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import shutil
import statistics
import subprocess
import time
from collections import defaultdict
from pathlib import Path


DATASETS = ("set_01", "set_02", "set_03")
PROFILES = ("weekday", "thu7", "chunhat")
HOURS = tuple(range(6, 13))
ALL_PROFILE_HOURS = tuple(range(6, 19))
SEEDS = (1, 2, 3, 4, 5)
ITERATIONS = 9_000
SEGMENT_ITERATIONS = 300
TIME_LIMIT_SECONDS = 1_800


def source_stem(dataset: str, profile: str) -> str:
    return f"hanoi_15x15_100_{dataset}_{profile}"


def policy_stem(dataset: str, policy: str) -> str:
    return f"hanoi_15x15_100_{dataset}_{policy.lower()}"


def read_vijl(path: Path) -> dict[tuple[int, int, int], float]:
    values: dict[tuple[int, int, int], float] = {}
    with path.open(encoding="utf-8") as stream:
        for line in stream:
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            i, j, hour, speed = line.split()[:4]
            values[(int(i), int(j), int(hour))] = float(speed)
    return values


def write_vijl(path: Path, values: dict[tuple[int, int, int], float], note: str) -> None:
    with path.open("w", encoding="utf-8") as stream:
        stream.write("# i j hour v_ijl_kph\n")
        stream.write(f"# {note}\n")
        for (i, j, hour), speed in sorted(values.items()):
            stream.write(f"{i} {j} {hour} {speed:.6f}\n")


def build_policy_profiles(source_root: Path, output_root: Path) -> None:
    for dataset in DATASETS:
        profiles = {
            profile: read_vijl(
                source_root / dataset / profile / f"{source_stem(dataset, profile)}.v_ijl_kph.txt"
            )
            for profile in PROFILES
        }
        keys = set(profiles[PROFILES[0]])
        if any(set(values) != keys for values in profiles.values()):
            raise ValueError(f"Traffic keys differ across profiles for {dataset}")

        all_speeds = [value for values in profiles.values() for value in values.values()]
        global_static_speed = statistics.fmean(all_speeds)
        p0 = {key: global_static_speed for key in keys}
        p1 = {
            key: statistics.fmean(profiles[profile][key] for profile in PROFILES)
            for key in keys
        }

        base_dir = source_root / dataset / "weekday"
        base_stem = source_stem(dataset, "weekday")
        for policy, values, note in (
            ("P0", p0, f"global static mean speed={global_static_speed:.6f} kph"),
            ("P1", p1, "edge/hour arithmetic mean across weekday, Saturday, Sunday"),
        ):
            directory = output_root / policy / dataset
            directory.mkdir(parents=True, exist_ok=True)
            stem = policy_stem(dataset, policy)
            for suffix in (".txt", ".truck_distance_m.txt", ".drone_distance_m.txt"):
                shutil.copy2(base_dir / f"{base_stem}{suffix}", directory / f"{stem}{suffix}")
            write_vijl(directory / f"{stem}.v_ijl_kph.txt", values, note)
            (directory / "metadata.json").write_text(
                json.dumps(
                    {
                        "policy": policy,
                        "dataset": dataset,
                        "static_speed_kph": global_static_speed if policy == "P0" else None,
                        "profiles_averaged": list(PROFILES),
                        "hours": list(ALL_PROFILE_HOURS),
                    },
                    indent=2,
                ) + "\n",
                encoding="utf-8",
            )


def policy_tasks() -> list[dict[str, object]]:
    tasks = []
    for dataset in DATASETS:
        for seed in SEEDS:
            tasks.append({"task_id": f"P0-{dataset}-seed{seed}", "policy": "P0", "dataset": dataset, "seed": seed})
        for hour in HOURS:
            for seed in SEEDS:
                tasks.append(
                    {
                        "task_id": f"P1-{dataset}-{hour}h-seed{seed}",
                        "policy": "P1", "dataset": dataset, "start_hour": hour, "seed": seed,
                    }
                )
    return tasks


def single_trip_tasks() -> list[dict[str, object]]:
    return [
        {
            "task_id": f"ST-{dataset}-{profile}-{hour}h-seed{seed}",
            "dataset": dataset, "profile": profile, "start_hour": hour, "seed": seed,
        }
        for dataset in DATASETS for profile in PROFILES for hour in HOURS for seed in SEEDS
    ]


def matrix(command: str, group_size: int) -> dict[str, list[dict[str, object]]]:
    tasks = policy_tasks() if command == "policy" else single_trip_tasks()
    include = []
    for index in range(0, len(tasks), group_size):
        batch = tasks[index:index + group_size]
        include.append(
            {
                "batch_id": f"{command}-batch{index // group_size + 1:03d}",
                "task_count": len(batch),
                "tasks_json": json.dumps(batch, separators=(",", ":")),
            }
        )
    return {"include": include}


def capture(pattern: str, text: str) -> str:
    matches = re.findall(pattern, text, flags=re.IGNORECASE | re.MULTILINE)
    return matches[-1] if matches else ""


def parse_solution(path: Path) -> dict[str, object]:
    text = path.read_text(encoding="utf-8", errors="replace") if path.is_file() else ""
    makespan = capture(r"Improved solution cost:\s*([-+0-9.eE]+)", text)
    feasibility = capture(r"Final solution feasibility:\s*(FEASIBLE|INFEASIBLE)", text)
    truck_lines = [line for line in text.splitlines() if re.match(r"^Truck \d+:", line)]
    drone_lines = [line for line in text.splitlines() if re.match(r"^Drone \d+:", line)]
    truck_trips = 0
    later_trip_customers = 0
    truck_customers = 0
    for line in truck_lines:
        route = [int(value) for value in line.split(":", 1)[1].split("|", 1)[0].split()]
        trip_index = 0
        in_trip = False
        for node in route:
            if node == 0:
                in_trip = False
            else:
                if not in_trip:
                    trip_index += 1
                    truck_trips += 1
                    in_trip = True
                truck_customers += 1
                if trip_index >= 2:
                    later_trip_customers += 1
    drone_customers = sum(
        value != "0"
        for line in drone_lines
        for value in line.split(":", 1)[1].split("|", 1)[0].split()
    )
    return {
        "improved_makespan_s": float(makespan) if makespan else None,
        "feasibility": feasibility or "UNKNOWN",
        "truck_trips": truck_trips,
        "depot_returns": max(0, truck_trips - len(truck_lines)),
        "later_trip_customers": later_trip_customers,
        "later_trip_customers_pct": 100.0 * later_trip_customers / 100.0,
        "truck_customers": truck_customers,
        "drone_customers": drone_customers,
        "drone_customers_pct": 100.0 * drone_customers / 100.0,
        "final_solution_text": text,
    }


def run_command(command: list[str], cwd: Path, log_path: Path, allowed=(0,)) -> tuple[int, float, str]:
    started = time.monotonic()
    completed = subprocess.run(command, cwd=cwd, text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    elapsed = time.monotonic() - started
    log_path.write_text("Command: " + " ".join(command) + "\n\n" + completed.stdout, encoding="utf-8")
    if completed.returncode not in allowed:
        raise RuntimeError(f"Command failed ({completed.returncode}): {log_path}")
    return completed.returncode, elapsed, completed.stdout


def solver_command(binary: Path, stem_path: Path, start_hour: int, seed: int) -> list[str]:
    return [
        str(binary), str(stem_path.with_suffix(".txt")),
        f"--truck-distance-file={stem_path}.truck_distance_m.txt",
        f"--drone-distance-file={stem_path}.drone_distance_m.txt",
        f"--truck-vijl-file={stem_path}.v_ijl_kph.txt",
        f"--start-hour={start_hour}", "--attempts=1", f"--iters={ITERATIONS}",
        f"--segment-iters={SEGMENT_ITERATIONS}", f"--time-limit={TIME_LIMIT_SECONDS}",
        f"--seed={seed}", "--auto-tune",
    ]


def actual_stem(source_root: Path, dataset: str, profile: str) -> Path:
    return source_root / dataset / profile / source_stem(dataset, profile)


def evaluate_plan(
    binary: Path, source_root: Path, dataset: str, actual_profile: str,
    start_hour: int, solution: Path, output_dir: Path, policy: str, seed: int,
) -> dict[str, object]:
    stem = actual_stem(source_root, dataset, actual_profile)
    output_dir.mkdir(parents=True, exist_ok=True)
    command = solver_command(binary, stem, start_hour, seed)
    command = [arg for arg in command if not arg.startswith(("--attempts=", "--iters=", "--segment-iters=", "--time-limit=", "--seed=")) and arg != "--auto-tune"]
    command.append(f"--evaluate-solution={solution}")
    return_code, elapsed, stdout = run_command(command, output_dir, output_dir / "evaluation.log", allowed=(0, 2))
    evaluated = output_dir / "output_solution_evaluated.txt"
    parsed = parse_solution(evaluated)
    validation = re.search(
        r"Total validation:\s*Makespan=([-+0-9.eE]+),\s*Deadline violation=([-+0-9.eE]+),\s*Energy violation=([-+0-9.eE]+),\s*Capacity violation=([-+0-9.eE]+)",
        stdout,
    )
    return {
        "policy": policy, "dataset": dataset, "actual_profile": actual_profile,
        "start_hour": start_hour, "seed": seed, "return_code": return_code,
        "feasibility": parsed["feasibility"], "realized_makespan_s": parsed["improved_makespan_s"],
        "deadline_violation": float(validation.group(2)) if validation else None,
        "energy_violation": float(validation.group(3)) if validation else None,
        "capacity_violation": float(validation.group(4)) if validation else None,
        "evaluation_wall_time_s": elapsed,
    }


def write_rows(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        return
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def run_policy_batch(binary: Path, source_root: Path, policy_root: Path, tasks: list[dict[str, object]], output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    optimization_rows, evaluation_rows = [], []
    for task in tasks:
        policy, dataset, seed = str(task["policy"]), str(task["dataset"]), int(task["seed"])
        planned_hour = 6 if policy == "P0" else int(task["start_hour"])
        task_dir = output / str(task["task_id"])
        task_dir.mkdir(parents=True, exist_ok=True)
        stem = policy_root / policy / dataset / policy_stem(dataset, policy)
        command = solver_command(binary, stem, planned_hour, seed)
        _, elapsed, _ = run_command(command, task_dir, task_dir / "solver.log")
        solution = task_dir / "output_solution_best.txt"
        saved_solution = task_dir / f"final_solution_{task['task_id']}.txt"
        shutil.copy2(solution, saved_solution)
        parsed = parse_solution(saved_solution)
        optimization_rows.append(
            {
                "task_id": task["task_id"], "policy": policy, "dataset": dataset,
                "planned_start_hour": planned_hour, "seed": seed,
                "feasibility": parsed["feasibility"], "planned_makespan_s": parsed["improved_makespan_s"],
                "wall_time_s": elapsed, "solution_file": saved_solution.name,
            }
        )
        evaluation_hours = HOURS if policy == "P0" else (planned_hour,)
        for profile in PROFILES:
            for hour in evaluation_hours:
                evaluation_rows.append(
                    evaluate_plan(binary, source_root, dataset, profile, hour, saved_solution,
                                  task_dir / "evaluations" / profile / f"{hour}h", policy, seed)
                )
    write_rows(output / "optimization_results.csv", optimization_rows)
    write_rows(output / "evaluation_results.csv", evaluation_rows)


def run_st_batch(binary: Path, source_root: Path, tasks: list[dict[str, object]], output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    rows = []
    for task in tasks:
        dataset, profile = str(task["dataset"]), str(task["profile"])
        hour, seed = int(task["start_hour"]), int(task["seed"])
        task_dir = output / str(task["task_id"])
        task_dir.mkdir(parents=True, exist_ok=True)
        command = solver_command(binary, actual_stem(source_root, dataset, profile), hour, seed)
        command.append("--truck-single-trip")
        _, elapsed, _ = run_command(command, task_dir, task_dir / "solver.log")
        solution = task_dir / "output_solution_best.txt"
        saved_solution = task_dir / f"final_solution_{task['task_id']}.txt"
        shutil.copy2(solution, saved_solution)
        parsed = parse_solution(saved_solution)
        rows.append(
            {
                "task_id": task["task_id"], "variant": "ST", "dataset": dataset,
                "profile": profile, "start_hour": hour, "seed": seed,
                "feasibility": parsed["feasibility"], "makespan_s": parsed["improved_makespan_s"],
                "truck_trips": parsed["truck_trips"], "depot_returns": parsed["depot_returns"],
                "later_trip_customers": parsed["later_trip_customers"],
                "drone_customers_pct": parsed["drone_customers_pct"],
                "wall_time_s": elapsed, "solution_file": saved_solution.name,
            }
        )
    write_rows(output / "single_trip_results.csv", rows)


def aggregate(input_root: Path, output: Path, experiment: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    patterns = (
        ("Optimization", "optimization_results.csv"),
        ("Evaluations", "evaluation_results.csv"),
    ) if experiment == "policy" else (("Single_Trip", "single_trip_results.csv"),)
    workbook = Workbook()
    workbook.remove(workbook.active)
    collected: dict[str, list[dict[str, str]]] = {}
    total_rows = 0
    for sheet_name, filename in patterns:
        rows = []
        for path in input_root.rglob(filename):
            with path.open(encoding="utf-8", newline="") as stream:
                rows.extend(csv.DictReader(stream))
        rows.sort(key=lambda row: tuple(row.get(key, "") for key in ("policy", "dataset", "actual_profile", "profile", "start_hour", "seed")))
        collected[sheet_name] = rows
        total_rows += len(rows)
        sheet = workbook.create_sheet(sheet_name)
        if not rows:
            sheet.append(["No data"])
            continue
        headers = list(rows[0])
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    expected = {"Optimization": 120, "Evaluations": 630} if experiment == "policy" else {"Single_Trip": 315}
    summary = workbook.create_sheet("Summary", 0)
    summary.append(["Experiment", experiment])
    summary.append(["Sheet", "Rows found", "Rows expected", "Status"])
    complete = True
    for sheet_name, expected_rows in expected.items():
        found = len(collected.get(sheet_name, []))
        complete &= found == expected_rows
        summary.append([sheet_name, found, expected_rows, "COMPLETE" if found == expected_rows else "INCOMPLETE"])
    summary.append([])
    summary.append(["Group", "Runs", "Feasible", "Infeasible", "Feasible rate (%)", "Mean makespan (s)"])
    grouped: dict[tuple[str, ...], list[dict[str, str]]] = defaultdict(list)
    detail_rows = collected.get("Evaluations" if experiment == "policy" else "Single_Trip", [])
    for row in detail_rows:
        key = (
            (row.get("policy", "ST")),
            row.get("actual_profile", row.get("profile", "")),
            row.get("start_hour", ""),
        )
        grouped[key].append(row)
    for key, rows in sorted(grouped.items()):
        feasible = sum(row.get("feasibility") == "FEASIBLE" for row in rows)
        makespan_key = "realized_makespan_s" if experiment == "policy" else "makespan_s"
        makespans = [float(row[makespan_key]) for row in rows if row.get(makespan_key)]
        summary.append([
            " / ".join(key), len(rows), feasible, len(rows) - feasible,
            100.0 * feasible / len(rows), statistics.fmean(makespans) if makespans else "",
        ])
    for cell in summary[2]:
        cell.font = Font(bold=True)
    summary.freeze_panes = "A3"

    solutions = workbook.create_sheet("Final_Solutions")
    solutions.append(["Solution file", "Artifact path", "Final solution text"])
    solution_paths = sorted(input_root.rglob("final_solution_*.txt"))
    for path in solution_paths:
        text = path.read_text(encoding="utf-8", errors="replace")
        solutions.append([path.name, str(path.relative_to(input_root)), text[:32767]])
    for cell in solutions[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    solutions.freeze_panes = "A2"
    solutions.column_dimensions["A"].width = 45
    solutions.column_dimensions["B"].width = 70
    solutions.column_dimensions["C"].width = 100

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(json.dumps({
        "experiment": experiment, "rows": total_rows, "solutions": len(solution_paths),
        "complete": complete, "workbook": str(output),
    }))


def main() -> None:
    parser = argparse.ArgumentParser()
    subparsers = parser.add_subparsers(dest="command", required=True)
    build = subparsers.add_parser("build-profiles")
    build.add_argument("--source-root", type=Path, required=True)
    build.add_argument("--output-root", type=Path, required=True)
    make_matrix = subparsers.add_parser("matrix")
    make_matrix.add_argument("--experiment", choices=("policy", "single-trip"), required=True)
    make_matrix.add_argument("--group-size", type=int, default=1)
    run_policy = subparsers.add_parser("run-policy")
    run_policy.add_argument("--binary", type=Path, required=True)
    run_policy.add_argument("--source-root", type=Path, required=True)
    run_policy.add_argument("--policy-root", type=Path, required=True)
    run_policy.add_argument("--tasks-json", required=True)
    run_policy.add_argument("--output", type=Path, required=True)
    run_st = subparsers.add_parser("run-single-trip")
    run_st.add_argument("--binary", type=Path, required=True)
    run_st.add_argument("--source-root", type=Path, required=True)
    run_st.add_argument("--tasks-json", required=True)
    run_st.add_argument("--output", type=Path, required=True)
    aggregate_parser = subparsers.add_parser("aggregate")
    aggregate_parser.add_argument("--input-root", type=Path, required=True)
    aggregate_parser.add_argument("--output", type=Path, required=True)
    aggregate_parser.add_argument("--experiment", choices=("policy", "single-trip"), required=True)
    args = parser.parse_args()

    if args.command == "build-profiles":
        build_policy_profiles(args.source_root, args.output_root)
    elif args.command == "matrix":
        print(json.dumps(matrix(args.experiment, args.group_size), separators=(",", ":")))
    elif args.command == "run-policy":
        run_policy_batch(args.binary.resolve(), args.source_root.resolve(), args.policy_root.resolve(), json.loads(args.tasks_json), args.output.resolve())
    elif args.command == "run-single-trip":
        run_st_batch(args.binary.resolve(), args.source_root.resolve(), json.loads(args.tasks_json), args.output.resolve())
    else:
        aggregate(args.input_root, args.output, args.experiment)


if __name__ == "__main__":
    main()
