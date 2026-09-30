#!/usr/bin/env python3
import argparse
import csv
import json
import re
import statistics
import zipfile
from collections import defaultdict
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


EXPECTED_N = (50, 100, 200)
EXPECTED_SPATIAL = (10, 20, 30, 40)
EXPECTED_INSTANCE_INDEX = (1, 3)
EXPECTED_SEGMENT_FACTORS = (1, 3, 5)
EXPECTED_RUNS = tuple(range(1, 6))
SEEDS = (1001, 2002, 3003, 4004, 5005)
ROUTE_RE = re.compile(
    r"^(Truck|Drone) (\d+):\s*(.*?)\|(Truck|Drone) Time:\s*([^|]+)\|"
    r"([^,]+),([^,]+),([^,]+),([^,]+)$"
)


def expected_task_ids() -> set[str]:
    return {
        f"{n}.{spatial}.{instance_index}-L{factor}n-run{run}"
        for n in EXPECTED_N
        for spatial in EXPECTED_SPATIAL
        for instance_index in EXPECTED_INSTANCE_INDEX
        for factor in EXPECTED_SEGMENT_FACTORS
        for run in EXPECTED_RUNS
    }


def as_float(value: object) -> Optional[float]:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def numeric_sort(value: object) -> tuple[int, ...]:
    return tuple(int(part) for part in str(value).split("."))


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    if not rows:
        return
    with path.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def read_results(results_dir: Path) -> tuple[list[dict[str, str]], dict[str, Path]]:
    rows: list[dict[str, str]] = []
    result_dirs: dict[str, Path] = {}
    for path in sorted(results_dir.rglob("result_*.csv")):
        with path.open(newline="", encoding="utf-8") as stream:
            for row in csv.DictReader(stream):
                rows.append(row)
                result_dirs[row["task_id"]] = path.parent
    rows.sort(
        key=lambda row: (
            int(row["n"]), int(row["spatial_index"]), int(row["instance_index"]),
            int(row["segment_factor"]), int(row["run"]),
        )
    )
    return rows, result_dirs


def build_instance_summary(rows: list[dict[str, str]]) -> list[dict[str, object]]:
    groups: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        groups[(row["instance"], row["segment_factor"])].append(row)

    references: dict[str, float] = {}
    group_costs: dict[tuple[str, str], list[float]] = {}
    for key, group in groups.items():
        costs = [
            value for row in group
            if row.get("feasible") == "FEASIBLE"
            and (value := as_float(row.get("improved_makespan_s"))) is not None
        ]
        group_costs[key] = costs
        if costs:
            instance = key[0]
            references[instance] = min(references.get(instance, float("inf")), min(costs))

    summary: list[dict[str, object]] = []
    for (instance, factor), group in groups.items():
        costs = group_costs[(instance, factor)]
        feasible_rows = [
            row for row in group
            if row.get("feasible") == "FEASIBLE"
            and as_float(row.get("improved_makespan_s")) is not None
        ]
        best_row = min(
            feasible_rows,
            key=lambda row: as_float(row["improved_makespan_s"]),
        ) if feasible_rows else None
        cpu_times = [value for row in group if (value := as_float(row.get("cpu_time_sec"))) is not None]
        wall_times = [value for row in group if (value := as_float(row.get("wall_runtime_sec"))) is not None]
        executed = [value for row in group if (value := as_float(row.get("executed_iterations"))) is not None]
        reference = references.get(instance)
        best = min(costs) if costs else None
        mean = statistics.fmean(costs) if costs else None
        sample_std = statistics.stdev(costs) if len(costs) >= 2 else (0.0 if costs else None)
        cv = 100.0 * sample_std / mean if sample_std is not None and mean else None
        best_gap = 100.0 * (best - reference) / reference if best is not None and reference else None
        mean_gap = 100.0 * (mean - reference) / reference if mean is not None and reference else None
        first = group[0]
        summary.append(
            {
                "n": int(first["n"]),
                "instance": instance,
                "segment_factor": int(factor),
                "segment_label": f"{factor}n",
                "segment_iterations": int(first["segment_iterations"]),
                "runs": len(group),
                "successful_runs": sum(row.get("status") == "SUCCESS" for row in group),
                "feasible_runs": len(costs),
                "C_best_s": best,
                "best_task_id": best_row["task_id"] if best_row else "",
                "best_run": int(best_row["run"]) if best_row else None,
                "best_seed": int(best_row["seed"]) if best_row else None,
                "C_mean_s": mean,
                "C_ref_s": reference,
                "best_gap_pct": best_gap,
                "mean_gap_pct": mean_gap,
                "sample_std_s": sample_std,
                "cv_pct": cv,
                "mean_cpu_time_s": statistics.fmean(cpu_times) if cpu_times else None,
                "mean_wall_time_s": statistics.fmean(wall_times) if wall_times else None,
                "mean_executed_iterations": statistics.fmean(executed) if executed else None,
                "iteration_limit_runs": sum(row.get("termination_reason") == "ITERATION_LIMIT" for row in group),
                "time_limit_runs": sum(row.get("termination_reason") == "TIME_LIMIT" for row in group),
            }
        )
    return sorted(summary, key=lambda row: (row["n"], numeric_sort(row["instance"]), row["segment_factor"]))


def build_paper_summary(
    rows: list[dict[str, str]], instance_summary: list[dict[str, object]]
) -> list[dict[str, object]]:
    output: list[dict[str, object]] = []
    for factor in EXPECTED_SEGMENT_FACTORS:
        instance_rows = [row for row in instance_summary if row["segment_factor"] == factor]
        best_gaps = [row["best_gap_pct"] for row in instance_rows if row["best_gap_pct"] is not None]
        mean_gaps = [row["mean_gap_pct"] for row in instance_rows if row["mean_gap_pct"] is not None]
        cvs = [row["cv_pct"] for row in instance_rows if row["cv_pct"] is not None]
        run_rows = [row for row in rows if int(row["segment_factor"]) == factor]
        cpu_times = [value for row in run_rows if (value := as_float(row.get("cpu_time_sec"))) is not None]
        output.append(
            {
                "L_seg": f"{factor}n",
                "best_gap_pct": statistics.fmean(best_gaps) if best_gaps else None,
                "mean_gap_pct": statistics.fmean(mean_gaps) if mean_gaps else None,
                "cv_pct": statistics.fmean(cvs) if cvs else None,
                "cpu_time_s": statistics.fmean(cpu_times) if cpu_times else None,
                "instances_in_gap": len(best_gaps),
                "runs": len(run_rows),
                "feasible_runs": sum(row.get("feasible") == "FEASIBLE" for row in run_rows),
                "iteration_limit_runs": sum(row.get("termination_reason") == "ITERATION_LIMIT" for row in run_rows),
                "time_limit_runs": sum(row.get("termination_reason") == "TIME_LIMIT" for row in run_rows),
            }
        )
    return output


def collect_solutions(
    rows: list[dict[str, str]], result_dirs: dict[str, Path]
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[str]]:
    solutions: list[dict[str, object]] = []
    routes: list[dict[str, object]] = []
    missing: list[str] = []
    for row in rows:
        filename = row.get("solution_file", "")
        path = result_dirs.get(row["task_id"], Path()) / filename if filename else None
        if path is None or not path.is_file():
            missing.append(row["task_id"])
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        chunks = [text[index : index + 30000] for index in range(0, len(text), 30000)] or [""]
        for part, chunk in enumerate(chunks, start=1):
            solutions.append(
                {
                    "task_id": row["task_id"],
                    "instance": row["instance"],
                    "segment_label": f"{row['segment_factor']}n",
                    "run": int(row["run"]),
                    "seed": int(row["seed"]),
                    "feasible": row["feasible"],
                    "makespan_s": as_float(row.get("improved_makespan_s")),
                    "part": part,
                    "parts": len(chunks),
                    "final_solution_text": chunk,
                }
            )
        for line in text.splitlines():
            match = ROUTE_RE.match(line.strip())
            if not match:
                continue
            routes.append(
                {
                    "task_id": row["task_id"],
                    "instance": row["instance"],
                    "segment_label": f"{row['segment_factor']}n",
                    "run": int(row["run"]),
                    "seed": int(row["seed"]),
                    "vehicle_type": match.group(1),
                    "vehicle_index": int(match.group(2)),
                    "route": match.group(3).strip(),
                    "cached_route_time_s": as_float(match.group(5).strip()),
                    "validated_route_time_s": as_float(match.group(6).strip()),
                    "deadline_violation": as_float(match.group(7).strip()),
                    "energy_violation": as_float(match.group(8).strip()),
                    "capacity_violation": as_float(match.group(9).strip()),
                }
            )
    return solutions, routes, missing


def write_solution_archive(
    path: Path, rows: list[dict[str, str]], result_dirs: dict[str, Path]
) -> int:
    archived = 0
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for row in rows:
            filename = row.get("solution_file", "")
            source = result_dirs.get(row["task_id"], Path()) / filename if filename else None
            if source is None or not source.is_file():
                continue
            archive.write(source, arcname=f"final_solutions/{filename}")
            archived += 1
    return archived


NUMBER_RE = re.compile(r"[-+]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][-+]?\d+)?")


def excel_value(value: object) -> object:
    if not isinstance(value, str):
        return value
    stripped = value.strip()
    if not NUMBER_RE.fullmatch(stripped):
        return value
    if re.fullmatch(r"[-+]?\d+", stripped):
        return int(stripped)
    return float(stripped)


def add_sheet(workbook: Workbook, name: str, rows: list[dict[str, object]]) -> None:
    sheet = workbook.create_sheet(name)
    if not rows:
        sheet.append(["No data"])
        return
    headers = list(rows[0])
    sheet.append(headers)
    for row in rows:
        sheet.append([excel_value(row.get(header)) for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        values = [str(header)] + [str(row.get(header, "") or "") for row in rows[:200]]
        width = min(max(max(len(value) for value in values) + 2, 12), 60)
        if header == "final_solution_text":
            width = 100
            for cell in sheet[get_column_letter(index)][1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        elif header == "route":
            width = 70
        sheet.column_dimensions[get_column_letter(index)].width = width


def build_workbook(
    path: Path,
    paper_summary: list[dict[str, object]],
    instance_summary: list[dict[str, object]],
    all_runs: list[dict[str, str]],
    solutions: list[dict[str, object]],
    routes: list[dict[str, object]],
    validation_rows: list[dict[str, object]],
    baseline_run_id: str,
    current_run_id: str,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_sheet(workbook, "Paper_Summary", paper_summary)
    add_sheet(workbook, "Instance_Summary", instance_summary)
    add_sheet(workbook, "All_Runs", all_runs)
    add_sheet(workbook, "Final_Solutions", solutions)
    add_sheet(workbook, "Routes", routes)
    limits_by_n: dict[int, set[int]] = defaultdict(set)
    for row in all_runs:
        try:
            limits_by_n[int(row["n"])].add(int(row["time_limit_sec"]))
        except (KeyError, TypeError, ValueError):
            continue
    time_limit_summary = "; ".join(
        f"n={n}:{','.join(str(value) for value in sorted(values))}"
        for n, values in sorted(limits_by_n.items())
    )
    stop_condition_summary = "; ".join(
        f"n={n}:IT_max or {','.join(str(value) for value in sorted(values))} seconds"
        for n, values in sorted(limits_by_n.items())
    )
    config = [
        {"parameter": "instances", "value": "{50,100,200}.{10,20,30,40}.{1,3}"},
        {"parameter": "baseline_run_id_n50_n100", "value": baseline_run_id},
        {"parameter": "current_run_id_n200", "value": current_run_id},
        {"parameter": "independent_runs_R", "value": len(SEEDS)},
        {"parameter": "common_seeds", "value": ",".join(map(str, SEEDS))},
        {"parameter": "solver_processes", "value": 360},
        {"parameter": "solver_jobs", "value": 204},
        {"parameter": "sequential_processes_per_job", "value": "n=50:5; n=100:2; n=200:1"},
        {"parameter": "process_time_limit_sec", "value": time_limit_summary},
        {"parameter": "solver_job_timeout_min", "value": 300},
        {"parameter": "process_stop_condition", "value": stop_condition_summary},
        {"parameter": "segment_lengths", "value": "n,3n,5n"},
        {"parameter": "iteration_budget", "value": "9*n*ceil(sqrt(n))"},
        {"parameter": "customer_demand_kg", "value": "(0,2]"},
        {"parameter": "truck_capacity_kg", "value": 1000},
        {"parameter": "truck_vmax_mph", "value": "edge-specific samples in [24,35]"},
        {"parameter": "theta_ijl", "value": "edge/time-specific values in [0.4,1]"},
        {"parameter": "drone_capacity_kg", "value": 2.27},
        {"parameter": "drone_energy_capacity_j", "value": 7200000},
        {"parameter": "drone_cruise_speed_mph", "value": 70},
        {"parameter": "drone_ascent_speed_mph", "value": 35},
        {"parameter": "drone_descent_speed_mph", "value": 17.5},
        {"parameter": "drone_altitude_m", "value": 50},
        {"parameter": "power_model", "value": "P(w)=24.2*w+1329.0 W"},
        {"parameter": "A2_gamma", "value": "gamma1=0.5,gamma2=0.3,gamma3=0.1,gamma4=0.3"},
        {"parameter": "adaptive_neighborhood_selection", "value": "enabled"},
        {"parameter": "mode_switching", "value": "disabled"},
        {"parameter": "simulated_annealing", "value": "disabled (T0=0)"},
        {"parameter": "destroy_and_repair", "value": "disabled"},
        {"parameter": "penalty_lambdas", "value": "fixed at 1 (kappa=0)"},
        {"parameter": "tabu_base", "value": "max(20,floor(2*sqrt(n)))"},
        {"parameter": "tabu_random_interval", "value": "[tau,2*tau], inclusive"},
    ]
    add_sheet(workbook, "Experiment_Config", config)
    add_sheet(workbook, "Validation", validation_rows)
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("results_dir", type=Path)
    parser.add_argument("--out-dir", type=Path, default=Path("v2_segment_aggregate"))
    parser.add_argument("--baseline-run-id", default="")
    parser.add_argument("--current-run-id", default="")
    parser.add_argument("--strict", action="store_true")
    args = parser.parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)

    rows, result_dirs = read_results(args.results_dir)
    expected = expected_task_ids()
    seen: dict[str, int] = defaultdict(int)
    for row in rows:
        seen[row.get("task_id", "")] += 1
    found = set(seen)
    missing = sorted(expected - found)
    unexpected = sorted(found - expected)
    duplicates = sorted(task_id for task_id, count in seen.items() if count > 1)
    failed = sorted(row["task_id"] for row in rows if row.get("status") != "SUCCESS")

    instance_summary = build_instance_summary(rows)
    paper_summary = build_paper_summary(rows, instance_summary)
    solutions, routes, missing_solutions = collect_solutions(rows, result_dirs)
    archived_solutions = write_solution_archive(
        args.out_dir / "v2_segment_final_solutions.zip", rows, result_dirs
    )
    incomplete_groups = [
        f"{row['instance']}:{row['segment_label']} (runs={row['runs']}, feasible={row['feasible_runs']})"
        for row in instance_summary
        if row["runs"] != len(EXPECTED_RUNS) or row["feasible_runs"] != len(EXPECTED_RUNS)
    ]

    validation = {
        "expected_runs": len(expected),
        "collected_rows": len(rows),
        "unique_task_ids": len(found),
        "successful_runs": sum(row.get("status") == "SUCCESS" for row in rows),
        "feasible_runs": sum(row.get("feasible") == "FEASIBLE" for row in rows),
        "final_solutions_collected": len({row["task_id"] for row in solutions}),
        "final_solutions_archived": archived_solutions,
        "route_rows_collected": len(routes),
        "missing_task_ids": missing,
        "unexpected_task_ids": unexpected,
        "duplicate_task_ids": duplicates,
        "failed_task_ids": failed,
        "missing_final_solutions": missing_solutions,
        "incomplete_or_infeasible_instance_segment_groups": incomplete_groups,
    }
    (args.out_dir / "v2_segment_validation.json").write_text(
        json.dumps(validation, indent=2) + "\n", encoding="utf-8"
    )

    write_csv(args.out_dir / "v2_segment_all_runs.csv", rows)
    write_csv(args.out_dir / "v2_segment_instance_summary.csv", instance_summary)
    write_csv(args.out_dir / "v2_segment_paper_summary.csv", paper_summary)

    validation_rows = []
    for key, value in validation.items():
        if isinstance(value, list):
            if value:
                validation_rows.extend({"check": key, "value": item} for item in value)
            else:
                validation_rows.append({"check": key, "value": "OK"})
        else:
            validation_rows.append({"check": key, "value": value})

    build_workbook(
        args.out_dir / "v2_segment_experiment.xlsx",
        paper_summary,
        instance_summary,
        rows,
        solutions,
        routes,
        validation_rows,
        args.baseline_run_id,
        args.current_run_id,
    )

    markdown = [
        "# V2 segment-length experiment",
        "",
        f"- Expected runs: {len(expected)}",
        f"- Collected rows: {len(rows)}",
        f"- Successful runs: {validation['successful_runs']}",
        f"- Feasible runs: {validation['feasible_runs']}",
        f"- Final solutions: {validation['final_solutions_collected']}",
        "",
        "| L_seg | Best gap (%) | Mean gap (%) | CV (%) | CPU time (s) |",
        "|---:|---:|---:|---:|---:|",
    ]
    for row in paper_summary:
        values = [row["best_gap_pct"], row["mean_gap_pct"], row["cv_pct"], row["cpu_time_s"]]
        formatted = [f"{value:.6f}" if value is not None else "N/A" for value in values]
        markdown.append(
            f"| {row['L_seg']} | {formatted[0]} | {formatted[1]} | {formatted[2]} | {formatted[3]} |"
        )
    (args.out_dir / "v2_segment_summary.md").write_text("\n".join(markdown) + "\n", encoding="utf-8")

    issues = bool(
        missing or unexpected or duplicates or failed or missing_solutions
        or archived_solutions != len(expected)
        or incomplete_groups or len(rows) != len(expected)
    )
    print(json.dumps(validation, indent=2))
    if args.strict and issues:
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
